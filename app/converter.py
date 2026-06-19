import os
import io
import uuid
import html as html_escape
import logging
from pathlib import Path
from typing import Optional, Tuple, List
from docx import Document as DocxDocument
from docx.oxml.ns import qn
from pdf2image import convert_from_path
from PIL import Image, ImageDraw, ImageFont

from app.config import get_settings

settings = get_settings()
logger = logging.getLogger(__name__)

ALLOWED_EXTENSIONS = {
    "docx": "document",
    "pdf": "pdf",
    "png": "image",
    "jpg": "image",
    "jpeg": "image",
    "gif": "image",
    "bmp": "image",
    "webp": "image",
}


def get_file_type(filename: str) -> Optional[str]:
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    return ALLOWED_EXTENSIONS.get(ext)


def is_allowed_file(filename: str) -> bool:
    return get_file_type(filename) is not None


def get_extension(filename: str) -> str:
    return filename.rsplit(".", 1)[-1].lower() if "." in filename else ""


def ensure_dirs():
    Path(settings.upload_dir).mkdir(parents=True, exist_ok=True)
    Path(settings.preview_dir).mkdir(parents=True, exist_ok=True)


def save_upload_file(file_content: bytes, original_filename: str) -> Tuple[str, str]:
    ensure_dirs()
    ext = get_extension(original_filename)
    unique_name = f"{uuid.uuid4().hex}.{ext}"
    file_path = os.path.join(settings.upload_dir, unique_name)
    with open(file_path, "wb") as f:
        f.write(file_content)
    return unique_name, file_path


def add_text_watermark_to_image(image: Image.Image, watermark_text: str, opacity: int = 80) -> Image.Image:
    if not watermark_text:
        return image
    image = image.convert("RGBA")
    width, height = image.size
    watermark = Image.new("RGBA", image.size, (255, 255, 255, 0))
    draw = ImageDraw.Draw(watermark)
    try:
        font_size = max(16, min(width, height) // 20)
        font = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", font_size)
    except (OSError, IOError):
        font = ImageFont.load_default()
    text_bbox = draw.textbbox((0, 0), watermark_text, font=font)
    text_width = text_bbox[2] - text_bbox[0]
    text_height = text_bbox[3] - text_bbox[1]
    spacing_x = text_width + width // 8
    spacing_y = text_height + height // 8
    for pos_x in range(0, width + spacing_x, spacing_x):
        for pos_y in range(0, height + spacing_y, spacing_y):
            draw.text((pos_x, pos_y), watermark_text, fill=(200, 200, 200, opacity), font=font)
    combined = Image.alpha_composite(image, watermark)
    return combined.convert("RGB")


def convert_docx_to_html(file_path: str, watermark_text: Optional[str] = None) -> Tuple[str, str]:
    doc = DocxDocument(file_path)
    html_parts = [
        "<!DOCTYPE html><html><head><meta charset='utf-8'><style>",
        "body { font-family: Arial, sans-serif; max-width: 800px; margin: 40px auto; padding: 20px; }",
        "h1 { color: #333; } h2 { color: #444; }",
        "table { border-collapse: collapse; margin: 20px 0; }",
        "table, th, td { border: 1px solid #ddd; }",
        "th, td { padding: 12px; text-align: left; }",
        "th { background-color: #f5f5f5; }",
        ".watermark {",
        "position: fixed; top: 50%; left: 50%;",
        "color: rgba(200, 200, 200, 0.3); font-size: 48px; font-weight: bold;",
        "pointer-events: none; z-index: 9999; white-space: nowrap;",
        "transform: translate(-50%, -50%) rotate(-45deg);",
        "}",
        "</style></head><body>",
    ]
    if watermark_text:
        html_parts.append(f"<div class='watermark'>{html_escape.escape(watermark_text)}</div>")
    for para in doc.paragraphs:
        style_name = para.style.name if para.style else ""
        text = html_escape.escape(para.text) if para.text else "&nbsp;"
        if style_name.startswith("Heading 1"):
            html_parts.append(f"<h1>{text}</h1>")
        elif style_name.startswith("Heading 2"):
            html_parts.append(f"<h2>{text}</h2>")
        elif style_name.startswith("Heading 3"):
            html_parts.append(f"<h3>{text}</h3>")
        else:
            html_parts.append(f"<p>{text}</p>")
    for table in doc.tables:
        html_parts.append("<table>")
        for i, row in enumerate(table.rows):
            html_parts.append("<tr>")
            for cell in row.cells:
                tag = "th" if i == 0 else "td"
                html_parts.append(f"<{tag}>{html_escape.escape(cell.text)}</{tag}>")
            html_parts.append("</tr>")
        html_parts.append("</table>")
    html_parts.append("</body></html>")
    html_content = "\n".join(html_parts)
    ensure_dirs()
    preview_filename = f"{uuid.uuid4().hex}.html"
    preview_path = os.path.join(settings.preview_dir, preview_filename)
    with open(preview_path, "w", encoding="utf-8") as f:
        f.write(html_content)
    return preview_filename, preview_path


def convert_pdf_to_images(file_path: str, watermark_text: Optional[str] = None) -> Tuple[str, str]:
    ensure_dirs()
    images = convert_from_path(file_path)
    preview_dir_name = f"{uuid.uuid4().hex}"
    preview_full_dir = os.path.join(settings.preview_dir, preview_dir_name)
    os.makedirs(preview_full_dir, exist_ok=True)
    for i, image in enumerate(images):
        if watermark_text:
            image = add_text_watermark_to_image(image, watermark_text)
        image_path = os.path.join(preview_full_dir, f"page_{i + 1}.png")
        image.save(image_path, "PNG")
    return preview_dir_name, preview_full_dir


def convert_image(file_path: str, watermark_text: Optional[str] = None) -> Tuple[str, str]:
    ensure_dirs()
    image = Image.open(file_path)
    if watermark_text:
        image = add_text_watermark_to_image(image, watermark_text)
    preview_filename = f"{uuid.uuid4().hex}.png"
    preview_path = os.path.join(settings.preview_dir, preview_filename)
    image.save(preview_path, "PNG")
    return preview_filename, preview_path


def convert_document(file_path: str, file_type: str, watermark_text: Optional[str] = None) -> Optional[Tuple[str, str]]:
    ext = get_extension(file_path)
    try:
        if ext == "docx" or file_type == "document":
            return convert_docx_to_html(file_path, watermark_text)
        elif ext == "pdf" or file_type == "pdf":
            return convert_pdf_to_images(file_path, watermark_text)
        elif file_type == "image":
            return convert_image(file_path, watermark_text)
    except Exception as e:
        logger.error(f"Conversion error: {e}")
        raise
    return None


def get_pdf_page_count(preview_dir_path: str) -> int:
    if not os.path.isdir(preview_dir_path):
        return 0
    return len([f for f in os.listdir(preview_dir_path) if f.endswith(".png") and f.startswith("page_")])
