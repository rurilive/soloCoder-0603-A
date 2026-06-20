import os
import uuid
import html as html_escape
import logging
import asyncio
from pathlib import Path
from typing import Optional, Tuple, Dict, Any
from concurrent.futures import ThreadPoolExecutor
from docx import Document as DocxDocument
from pdf2image import convert_from_path
from PIL import Image, ImageDraw, ImageFont

logger = logging.getLogger(__name__)

ALLOWED_EXTENSIONS = {
    "docx": "document",
    "pdf": "pdf",
    "png": "image",
    "jpg": "image",
    "jpeg": "image",
    "gif": "image",
    "bmp": "image",
    "webp": "image",
    "txt": "text",
    "md": "text",
    "csv": "text",
    "log": "text",
    "json": "text",
    "xml": "text",
    "yaml": "text",
    "yml": "text",
    "ini": "text",
    "cfg": "text",
    "conf": "text",
    "py": "text",
    "js": "text",
    "ts": "text",
    "html": "text",
    "css": "text",
}

_executor = ThreadPoolExecutor(max_workers=4)


def get_extension(filename: str) -> str:
    return filename.rsplit(".", 1)[-1].lower() if "." in filename else ""


def get_file_type(filename: str) -> Optional[str]:
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    return ALLOWED_EXTENSIONS.get(ext)


def is_allowed_file(filename: str) -> bool:
    return get_file_type(filename) is not None


def ensure_dirs(upload_dir: str, preview_dir: str):
    Path(upload_dir).mkdir(parents=True, exist_ok=True)
    Path(preview_dir).mkdir(parents=True, exist_ok=True)


def add_text_watermark_to_image(image: Image.Image, watermark_text: str, opacity: int = 80) -> Image.Image:
    if not watermark_text:
        return image
    image = image.convert("RGBA")
    width, height = image.size
    watermark = Image.new("RGBA", image.size, (255, 255, 255, 0))
    draw = ImageDraw.Draw(watermark)
    try:
        font_size = max(16, min(width, height) // 20)
        font = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", font_size)
    except (OSError, IOError):
        font = ImageFont.load_default()
    text_bbox = draw.textbbox((0, 0), watermark_text, font=font)
    text_width = text_bbox[2] - text_bbox[0]
    text_height = text_bbox[3] - text_bbox[1]
    spacing_x = text_width + width // 8
    spacing_y = text_height + height // 8
    for pos_x in range(0, width + spacing_x, spacing_x):
        for pos_y in range(0, height + spacing_y, spacing_y):
            draw.text((pos_x, pos_y), watermark_text, fill=(200, 200, 200, opacity), font=font)
    combined = Image.alpha_composite(image, watermark)
    return combined.convert("RGB")


def convert_docx_to_html(file_path: str, preview_dir: str, watermark_text: Optional[str] = None) -> Tuple[str, str, str]:
    doc = DocxDocument(file_path)
    html_parts = [
        "<!DOCTYPE html><html><head><meta charset='utf-8'><style>",
        "body { font-family: Arial, sans-serif; max-width: 800px; margin: 40px auto; padding: 20px; }",
        "h1 { color: #333; } h2 { color: #444; }",
        "table { border-collapse: collapse; margin: 20px 0; }",
        "table, th, td { border: 1px solid #ddd; }",
        "th, td { padding: 12px; text-align: left; }",
        "th { background-color: #f5f5f5; }",
        ".watermark {",
        "position: fixed; top: 50%; left: 50%;",
        "color: rgba(200, 200, 200, 0.3); font-size: 48px; font-weight: bold;",
        "pointer-events: none; z-index: 9999; white-space: nowrap;",
        "transform: translate(-50%, -50%) rotate(-45deg);",
        "}",
        "</style></head><body>",
    ]
    if watermark_text:
        html_parts.append(f"<div class='watermark'>{html_escape.escape(watermark_text)}</div>")
    for para in doc.paragraphs:
        style_name = para.style.name if para.style else ""
        text = html_escape.escape(para.text) if para.text else "&nbsp;"
        if style_name.startswith("Heading 1"):
            html_parts.append(f"<h1>{text}</h1>")
        elif style_name.startswith("Heading 2"):
            html_parts.append(f"<h2>{text}</h2>")
        elif style_name.startswith("Heading 3"):
            html_parts.append(f"<h3>{text}</h3>")
        else:
            html_parts.append(f"<p>{text}</p>")
    for table in doc.tables:
        html_parts.append("<table>")
        for i, row in enumerate(table.rows):
            html_parts.append("<tr>")
            for cell in row.cells:
                tag = "th" if i == 0 else "td"
                html_parts.append(f"<{tag}>{html_escape.escape(cell.text)}</{tag}>")
            html_parts.append("</tr>")
        html_parts.append("</table>")
    html_parts.append("</body></html>")
    html_content = "\n".join(html_parts)
    Path(preview_dir).mkdir(parents=True, exist_ok=True)
    preview_filename = f"{uuid.uuid4().hex}.html"
    preview_path = os.path.join(preview_dir, preview_filename)
    with open(preview_path, "w", encoding="utf-8") as f:
        f.write(html_content)
    return preview_filename, preview_path, "html"


def convert_pdf_to_images(file_path: str, preview_dir: str, watermark_text: Optional[str] = None, progress_callback=None) -> Tuple[str, str, str]:
    Path(preview_dir).mkdir(parents=True, exist_ok=True)
    images = convert_from_path(file_path)
    total = len(images)
    preview_dir_name = f"{uuid.uuid4().hex}"
    preview_full_dir = os.path.join(preview_dir, preview_dir_name)
    os.makedirs(preview_full_dir, exist_ok=True)
    for i, image in enumerate(images):
        if watermark_text:
            image = add_text_watermark_to_image(image, watermark_text)
        image_path = os.path.join(preview_full_dir, f"page_{i + 1}.png")
        image.save(image_path, "PNG")
        if progress_callback:
            progress_callback(int((i + 1) / total * 100))
    return preview_dir_name, preview_full_dir, "pdf_images"


def convert_image(file_path: str, preview_dir: str, watermark_text: Optional[str] = None) -> Tuple[str, str, str]:
    Path(preview_dir).mkdir(parents=True, exist_ok=True)
    image = Image.open(file_path)
    if watermark_text:
        image = add_text_watermark_to_image(image, watermark_text)
    preview_filename = f"{uuid.uuid4().hex}.png"
    preview_path = os.path.join(preview_dir, preview_filename)
    image.save(preview_path, "PNG")
    return preview_filename, preview_path, "image"


def convert_text_to_html(file_path: str, preview_dir: str, watermark_text: Optional[str] = None) -> Tuple[str, str, str]:
    with open(file_path, "r", encoding="utf-8", errors="replace") as f:
        raw_content = f.read()
    html_parts = [
        "<!DOCTYPE html><html><head><meta charset='utf-8'><style>",
        "body { font-family: 'Courier New', Courier, monospace; max-width: 960px; margin: 40px auto; padding: 20px; white-space: pre-wrap; word-wrap: break-word; font-size: 14px; line-height: 1.6; }",
        ".watermark {",
        "position: fixed; top: 50%; left: 50%;",
        "color: rgba(200, 200, 200, 0.3); font-size: 48px; font-weight: bold;",
        "pointer-events: none; z-index: 9999; white-space: nowrap;",
        "transform: translate(-50%, -50%) rotate(-45deg);",
        "}",
        "</style></head><body>",
    ]
    if watermark_text:
        html_parts.append(f"<div class='watermark'>{html_escape.escape(watermark_text)}</div>")
    html_parts.append(html_escape.escape(raw_content))
    html_parts.append("</body></html>")
    html_content = "\n".join(html_parts)
    Path(preview_dir).mkdir(parents=True, exist_ok=True)
    preview_filename = f"{uuid.uuid4().hex}.html"
    preview_path = os.path.join(preview_dir, preview_filename)
    with open(preview_path, "w", encoding="utf-8") as f:
        f.write(html_content)
    return preview_filename, preview_path, "html"


def convert_pdf_native(file_path: str, preview_dir: str) -> Tuple[str, str, str]:
    Path(preview_dir).mkdir(parents=True, exist_ok=True)
    import fitz
    doc = fitz.open(file_path)
    total_pages = doc.page_count
    doc.close()
    preview_filename = os.path.basename(file_path)
    dest_path = os.path.join(preview_dir, f"{uuid.uuid4().hex}.pdf")
    import shutil
    shutil.copy2(file_path, dest_path)
    return os.path.basename(dest_path), dest_path, "pdf_native"


def convert_document(
    file_path: str,
    file_type: str,
    preview_dir: str,
    watermark_text: Optional[str] = None,
    use_pdf_native: bool = False,
    progress_callback=None,
) -> Optional[Tuple[str, str, str]]:
    ext = get_extension(file_path)
    try:
        if ext == "docx" or file_type == "document":
            return convert_docx_to_html(file_path, preview_dir, watermark_text)
        elif ext == "pdf" or file_type == "pdf":
            if use_pdf_native:
                return convert_pdf_native(file_path, preview_dir)
            return convert_pdf_to_images(file_path, preview_dir, watermark_text, progress_callback)
        elif file_type == "image":
            return convert_image(file_path, preview_dir, watermark_text)
        elif file_type == "text":
            return convert_text_to_html(file_path, preview_dir, watermark_text)
    except Exception as e:
        logger.error(f"Conversion error: {e}")
        raise
    return None


def get_pdf_page_count(preview_dir_path: str) -> int:
    if not os.path.isdir(preview_dir_path):
        return 0
    return len([f for f in os.listdir(preview_dir_path) if f.endswith(".png") and f.startswith("page_")])


def search_pdf_text(file_path: str, query: str) -> list:
    import fitz
    doc = fitz.open(file_path)
    results = []
    for page_num in range(doc.page_count):
        page = doc[page_num]
        text_instances = page.search_for(query)
        for inst in text_instances:
            results.append({
                "page": page_num + 1,
                "text": query,
                "x0": inst.x0,
                "y0": inst.y0,
                "x1": inst.x1,
                "y1": inst.y1,
            })
    doc.close()
    return results


async def run_convert_async(
    file_path: str,
    file_type: str,
    preview_dir: str,
    watermark_text: Optional[str] = None,
    use_pdf_native: bool = False,
    progress_callback=None,
) -> Optional[Tuple[str, str, str]]:
    loop = asyncio.get_event_loop()
    return await loop.run_in_executor(
        _executor,
        convert_document,
        file_path,
        file_type,
        preview_dir,
        watermark_text,
        use_pdf_native,
        progress_callback,
    )


async def run_search_async(file_path: str, query: str) -> list:
    loop = asyncio.get_event_loop()
    return await loop.run_in_executor(_executor, search_pdf_text, file_path, query)
