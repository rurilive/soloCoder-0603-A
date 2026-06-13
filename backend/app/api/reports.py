from fastapi import APIRouter, Depends, HTTPException, status
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from sqlalchemy import func, cast, Date
from datetime import datetime
from io import BytesIO
from urllib.parse import quote
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from ..schemas import EventStatistics
from ..models import Event, Registration, User, CheckInRecord
from ..database import get_db
from ..dependencies import get_current_organizer

router = APIRouter(prefix="/reports", tags=["reports"])

HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
THIN_BORDER = Border(
    left=Side(style='thin'),
    right=Side(style='thin'),
    top=Side(style='thin'),
    bottom=Side(style='thin')
)
TITLE_FONT = Font(bold=True, size=14, color="1F4E78")

def _style_header_row(ws, row_num, num_cols):
    for col in range(1, num_cols + 1):
        cell = ws.cell(row=row_num, column=col)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = THIN_BORDER

def _style_data_cell(ws, row, col):
    cell = ws.cell(row=row, column=col)
    cell.border = THIN_BORDER
    cell.alignment = Alignment(vertical='center', wrap_text=True)

def _auto_width(ws, num_cols):
    for col in range(1, num_cols + 1):
        max_length = 0
        column_letter = ws.cell(row=1, column=col).column_letter
        for row in ws.iter_rows(min_col=col, max_col=col, values_only=True):
            for cell_value in row:
                if cell_value:
                    cell_length = len(str(cell_value))
                    if cell_length > max_length:
                        max_length = cell_length
        adjusted_width = min(max_length + 4, 50)
        ws.column_dimensions[column_letter].width = adjusted_width

def _get_event_and_verify(event_id: int, organizer: User, db: Session):
    event = db.query(Event).filter(Event.id == event_id).first()
    if not event:
        raise HTTPException(status_code=404, detail="Event not found")
    if event.organizer_id != organizer.id:
        raise HTTPException(status_code=403, detail="Not authorized to access this event's reports")
    return event

def _collect_form_fields(registrations):
    field_set = set()
    field_order = []
    for reg in registrations:
        if reg.form_data:
            for key in reg.form_data.keys():
                if key not in field_set:
                    field_set.add(key)
                    field_order.append(key)
    return field_order

def _export_workbook_to_bytes(wb: Workbook) -> BytesIO:
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer

def _format_dt(dt):
    if dt is None:
        return ""
    if isinstance(dt, datetime):
        return dt.strftime("%Y-%m-%d %H:%M:%S")
    return str(dt)

def _sanitize_filename_for_ascii(name: str) -> str:
    import re
    result = re.sub(r'[\\/:*?"<>|]', '_', name)
    ascii_name = result.encode('ascii', 'ignore').decode('ascii')
    ascii_name = re.sub(r'_+', '_', ascii_name).strip('_')
    if len(ascii_name) < 3 or not ascii_name.replace('_', '').replace('.', '').isalnum():
        return None
    return ascii_name

def _build_content_disposition(filename: str, ascii_fallback: str = None) -> str:
    utf8_encoded = quote(filename, safe='', encoding='utf-8')
    safe_ascii = _sanitize_filename_for_ascii(ascii_fallback or filename)
    if safe_ascii:
        return f"attachment; filename=\"{safe_ascii}\"; filename*=UTF-8''{utf8_encoded}"
    return f"attachment; filename*=UTF-8''{utf8_encoded}"

def _make_export_response(buffer: BytesIO, filename: str, ascii_fallback: str = None) -> StreamingResponse:
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": _build_content_disposition(filename, ascii_fallback),
            "Cache-Control": "no-cache, no-store, must-revalidate",
            "Pragma": "no-cache",
            "Expires": "0",
            "X-Suggested-Filename": quote(filename, safe='', encoding='utf-8'),
        }
    )

def _status_label(status):
    mapping = {
        "confirmed": "已确认",
        "waitlisted": "候补",
        "cancelled": "已取消",
        "pending_confirmation": "待确认"
    }
    return mapping.get(status, status)

def _build_event_info_sheet(wb: Workbook, event: Event, stats: dict):
    ws = wb.active
    ws.title = "活动概览"

    ws.merge_cells('A1:D1')
    title_cell = ws['A1']
    title_cell.value = event.title
    title_cell.font = TITLE_FONT
    title_cell.alignment = Alignment(horizontal='center', vertical='center')

    info_rows = [
        ("活动ID", event.id),
        ("活动描述", event.description or ""),
        ("开始时间", _format_dt(event.start_time)),
        ("结束时间", _format_dt(event.end_time)),
        ("活动地点", event.location),
        ("最大容量", event.max_capacity),
        ("活动状态", event.status),
        ("创建时间", _format_dt(event.created_at)),
        ("", ""),
        ("报名统计", ""),
        ("  总报名人数", stats["total_registrations"]),
        ("  已确认报名", stats["confirmed_registrations"]),
        ("  候补名单", stats["waitlisted_registrations"]),
        ("  待确认递补", stats["pending_confirmation_registrations"]),
        ("  已取消报名", stats["cancelled_registrations"]),
        ("", ""),
        ("签到统计", ""),
        ("  总签到人数", stats["total_checkins"]),
        ("  签到率", f"{stats['checkin_rate']:.1%}"),
        ("  报名率", f"{stats['registration_rate']:.1%}"),
        ("导出时间", _format_dt(datetime.utcnow()))
    ]

    for idx, (k, v) in enumerate(info_rows, start=3):
        ws.cell(row=idx, column=1, value=k).font = Font(bold=True) if k and not k.startswith("  ") else Font()
        ws.cell(row=idx, column=2, value=v)

    ws.column_dimensions['A'].width = 22
    ws.column_dimensions['B'].width = 50

    return ws

def _build_registrations_sheet(wb: Workbook, event: Event, db: Session, form_fields):
    ws = wb.create_sheet("报名数据")

    headers = [
        "报名ID", "用户ID", "用户名", "用户邮箱", "状态",
        "票号", "候补位置", "签到状态", "签到时间",
        "报名时间", "递补通知时间", "递补确认时间"
    ] + form_fields

    for col_idx, h in enumerate(headers, start=1):
        ws.cell(row=1, column=col_idx, value=h)
    _style_header_row(ws, 1, len(headers))

    query = db.query(Registration, User).outerjoin(
        User, Registration.user_id == User.id
    ).filter(Registration.event_id == event.id).order_by(Registration.created_at.asc())

    row_num = 2
    for reg, user in query.all():
        values = [
            reg.id,
            reg.user_id,
            user.username if user else "",
            user.email if user else "",
            _status_label(reg.status),
            reg.ticket_code or "",
            reg.waitlist_position or "",
            "已签到" if reg.check_in else "未签到",
            _format_dt(reg.check_in_time),
            _format_dt(reg.created_at),
            _format_dt(reg.waitlist_offer_sent_at),
            _format_dt(reg.waitlist_confirmed_at)
        ]
        for field in form_fields:
            values.append(reg.form_data.get(field, "") if reg.form_data else "")

        for col_idx, val in enumerate(values, start=1):
            ws.cell(row=row_num, column=col_idx, value=val)
            _style_data_cell(ws, row_num, col_idx)
        row_num += 1

    _auto_width(ws, len(headers))
    ws.freeze_panes = "A2"

def _build_checkin_sheet(wb: Workbook, event: Event, db: Session):
    ws = wb.create_sheet("签到明细")

    headers = [
        "签到记录ID", "报名ID", "用户名", "用户邮箱", "票号",
        "入口", "设备", "签到时间"
    ]

    for col_idx, h in enumerate(headers, start=1):
        ws.cell(row=1, column=col_idx, value=h)
    _style_header_row(ws, 1, len(headers))

    from ..models import Device
    query = db.query(CheckInRecord, Registration, User, Device).outerjoin(
        Registration, CheckInRecord.registration_id == Registration.id
    ).outerjoin(
        User, Registration.user_id == User.id
    ).outerjoin(
        Device, CheckInRecord.device_id == Device.id
    ).filter(CheckInRecord.event_id == event.id).order_by(CheckInRecord.check_in_time.asc())

    row_num = 2
    for record, reg, user, device in query.all():
        values = [
            record.id,
            record.registration_id,
            user.username if user else "",
            user.email if user else "",
            reg.ticket_code if reg else "",
            record.entrance,
            f"{device.name} ({device.device_id})" if device else "无设备",
            _format_dt(record.check_in_time)
        ]
        for col_idx, val in enumerate(values, start=1):
            ws.cell(row=row_num, column=col_idx, value=val)
            _style_data_cell(ws, row_num, col_idx)
        row_num += 1

    _auto_width(ws, len(headers))
    ws.freeze_panes = "A2"

def _build_hourly_distribution_sheet(wb: Workbook, stats: dict):
    hourly = stats.get("checkin_hourly_distribution", [])
    if not hourly:
        return

    ws = wb.create_sheet("时段签到分布")

    ws.merge_cells('A1:C1')
    title = ws['A1']
    title.value = "签到时段分布（按小时）"
    title.font = TITLE_FONT
    title.alignment = Alignment(horizontal='center')

    headers = ["时段", "签到人次", "占比"]
    for c, h in enumerate(headers, 1):
        ws.cell(row=3, column=c, value=h)
    _style_header_row(ws, 3, len(headers))

    total = sum(h["count"] for h in hourly) or 1
    row_num = 4
    for h in hourly:
        ws.cell(row=row_num, column=1, value=h["hour"])
        ws.cell(row=row_num, column=2, value=h["count"])
        ws.cell(row=row_num, column=3, value=f"{h['count'] / total:.1%}")
        for c in range(1, 4):
            _style_data_cell(ws, row_num, c)
        row_num += 1

    ws.column_dimensions['A'].width = 16
    ws.column_dimensions['B'].width = 14
    ws.column_dimensions['C'].width = 14
    ws.freeze_panes = "A4"

def _build_form_fields_sheet(wb: Workbook, stats: dict):
    form_stats = stats.get("form_field_stats", [])
    if not form_stats:
        return

    ws = wb.create_sheet("表单字段统计")

    ws.merge_cells('A1:E1')
    title = ws['A1']
    title.value = "报名表单字段统计"
    title.font = TITLE_FONT
    title.alignment = Alignment(horizontal='center')

    headers = ["字段名称", "总回答数", "已填写", "未填写", "填写率"]
    for c, h in enumerate(headers, 1):
        ws.cell(row=3, column=c, value=h)
    _style_header_row(ws, 3, len(headers))

    row_num = 4
    for field in form_stats:
        ws.cell(row=row_num, column=1, value=field["field_name"])
        ws.cell(row=row_num, column=2, value=field["total_responses"])
        ws.cell(row=row_num, column=3, value=field["filled_count"])
        ws.cell(row=row_num, column=4, value=field["empty_count"])
        ws.cell(row=row_num, column=5, value=f"{field['fill_rate']:.1%}")
        for c in range(1, 6):
            _style_data_cell(ws, row_num, c)
        row_num += 1

    row_num += 1
    ws.cell(row=row_num, column=1, value="热门值分布").font = Font(bold=True, size=11, color="1F4E78")
    row_num += 1

    for field in form_stats:
        if not field.get("top_values"):
            continue
        ws.cell(row=row_num, column=1, value=field["field_name"]).font = Font(bold=True)
        row_num += 1
        for tv in field["top_values"]:
            ws.cell(row=row_num, column=2, value=tv["value"])
            ws.cell(row=row_num, column=3, value=tv["count"])
            _style_data_cell(ws, row_num, 2)
            _style_data_cell(ws, row_num, 3)
            row_num += 1
        row_num += 1

    ws.column_dimensions['A'].width = 24
    ws.column_dimensions['B'].width = 14
    ws.column_dimensions['C'].width = 12
    ws.column_dimensions['D'].width = 12
    ws.column_dimensions['E'].width = 12
    ws.freeze_panes = "A4"

def _build_no_show_sheet(wb: Workbook, event: Event, db: Session):
    ws = wb.create_sheet("未签到名单")

    ws.merge_cells('A1:F1')
    title = ws['A1']
    title.value = "已报名未签到人员名单"
    title.font = TITLE_FONT
    title.alignment = Alignment(horizontal='center')

    headers = ["报名ID", "用户名", "用户邮箱", "票号", "报名时间", "状态"]
    for c, h in enumerate(headers, 1):
        ws.cell(row=3, column=c, value=h)
    _style_header_row(ws, 3, len(headers))

    query = db.query(Registration, User).outerjoin(
        User, Registration.user_id == User.id
    ).filter(
        Registration.event_id == event.id,
        Registration.status == "confirmed",
        Registration.check_in == False
    ).order_by(Registration.created_at.asc())

    row_num = 4
    for reg, user in query.all():
        values = [
            reg.id,
            user.username if user else "",
            user.email if user else "",
            reg.ticket_code or "",
            _format_dt(reg.created_at),
            _status_label(reg.status)
        ]
        for col_idx, val in enumerate(values, start=1):
            ws.cell(row=row_num, column=col_idx, value=val)
            _style_data_cell(ws, row_num, col_idx)
        row_num += 1

    ws.column_dimensions['A'].width = 12
    ws.column_dimensions['B'].width = 18
    ws.column_dimensions['C'].width = 28
    ws.column_dimensions['D'].width = 20
    ws.column_dimensions['E'].width = 22
    ws.column_dimensions['F'].width = 12
    ws.freeze_panes = "A4"

def _build_summary_sheet(wb: Workbook, event: Event, stats: dict, db: Session):
    ws = wb.create_sheet("统计摘要")

    ws.merge_cells('A1:F1')
    title = ws['A1']
    title.value = f"{event.title} - 统计分析摘要"
    title.font = TITLE_FONT
    title.alignment = Alignment(horizontal='center')

    row_ptr = 3

    ws.cell(row=row_ptr, column=1, value="一、核心指标").font = Font(bold=True, size=12, color="1F4E78")
    row_ptr += 1
    kpi_headers = ["指标", "数值"]
    for c, h in enumerate(kpi_headers, 1):
        ws.cell(row=row_ptr, column=c, value=h)
    _style_header_row(ws, row_ptr, len(kpi_headers))
    row_ptr += 1

    kpi_data = [
        ("总报名人数", stats["total_registrations"]),
        ("已确认", stats["confirmed_registrations"]),
        ("候补", stats["waitlisted_registrations"]),
        ("已取消", stats["cancelled_registrations"]),
        ("总签到数", stats["total_checkins"]),
        ("未签到数", stats.get("no_show_count", 0)),
        ("签到率", f"{stats['checkin_rate']:.1%}"),
        ("未签到率", f"{stats.get('no_show_rate', 0):.1%}"),
        ("报名率(相对容量)", f"{stats['registration_rate']:.1%}"),
    ]
    for k, v in kpi_data:
        ws.cell(row=row_ptr, column=1, value=k)
        ws.cell(row=row_ptr, column=2, value=v)
        _style_data_cell(ws, row_ptr, 1)
        _style_data_cell(ws, row_ptr, 2)
        row_ptr += 1
    row_ptr += 1

    ws.cell(row=row_ptr, column=1, value="二、入口签到分布").font = Font(bold=True, size=12, color="1F4E78")
    row_ptr += 1
    entrance_headers = ["入口名称", "签到人次", "占比"]
    for c, h in enumerate(entrance_headers, 1):
        ws.cell(row=row_ptr, column=c, value=h)
    _style_header_row(ws, row_ptr, len(entrance_headers))
    row_ptr += 1

    total = sum(stats["entrance_counts"].values()) or 1
    for entrance, count in stats["entrance_counts"].items():
        ws.cell(row=row_ptr, column=1, value=entrance)
        ws.cell(row=row_ptr, column=2, value=count)
        ws.cell(row=row_ptr, column=3, value=f"{count / total:.1%}")
        for c in range(1, 4):
            _style_data_cell(ws, row_ptr, c)
        row_ptr += 1
    row_ptr += 1

    ws.cell(row=row_ptr, column=1, value="三、报名状态分布").font = Font(bold=True, size=12, color="1F4E78")
    row_ptr += 1
    status_headers = ["状态", "人数", "占比"]
    for c, h in enumerate(status_headers, 1):
        ws.cell(row=row_ptr, column=c, value=h)
    _style_header_row(ws, row_ptr, len(status_headers))
    row_ptr += 1

    total_reg = stats["total_registrations"] or 1
    for status_key, count in stats["status_distribution"].items():
        ws.cell(row=row_ptr, column=1, value=_status_label(status_key))
        ws.cell(row=row_ptr, column=2, value=count)
        ws.cell(row=row_ptr, column=3, value=f"{count / total_reg:.1%}")
        for c in range(1, 4):
            _style_data_cell(ws, row_ptr, c)
        row_ptr += 1

    ws.column_dimensions['A'].width = 28
    ws.column_dimensions['B'].width = 18
    ws.column_dimensions['C'].width = 18

def _compute_form_field_stats(event_id: int, db: Session) -> list:
    registrations = db.query(Registration).filter(
        Registration.event_id == event_id,
        Registration.form_data.isnot(None)
    ).all()

    if not registrations:
        return []

    field_stats = {}
    for reg in registrations:
        if not reg.form_data:
            continue
        for field, value in reg.form_data.items():
            if field not in field_stats:
                field_stats[field] = {"total": 0, "values": {}, "empty": 0}
            field_stats[field]["total"] += 1
            if value is None or value == "":
                field_stats[field]["empty"] += 1
            else:
                val_str = str(value)
                field_stats[field]["values"][val_str] = field_stats[field]["values"].get(val_str, 0) + 1

    result = []
    for field, stats in field_stats.items():
        filled = stats["total"] - stats["empty"]
        top_values = sorted(stats["values"].items(), key=lambda x: x[1], reverse=True)[:5]
        result.append({
            "field_name": field,
            "total_responses": stats["total"],
            "filled_count": filled,
            "empty_count": stats["empty"],
            "fill_rate": filled / stats["total"] if stats["total"] > 0 else 0,
            "top_values": [{"value": v, "count": c} for v, c in top_values],
            "unique_count": len(stats["values"])
        })
    return result

def _compute_hourly_distribution(event_id: int, db: Session) -> list:
    from sqlalchemy import extract
    hourly_stats = db.query(
        extract('hour', CheckInRecord.check_in_time).label('hour'),
        func.count(CheckInRecord.id).label('count')
    ).filter(CheckInRecord.event_id == event_id).group_by(
        extract('hour', CheckInRecord.check_in_time)
    ).order_by(extract('hour', CheckInRecord.check_in_time)).all()

    result = []
    for hour, count in hourly_stats:
        result.append({
            "hour": f"{int(hour):02d}:00",
            "hour_num": int(hour),
            "count": count
        })
    return result

def _compute_statistics(event_id: int, db: Session) -> dict:
    event = db.query(Event).filter(Event.id == event_id).first()
    if not event:
        return {}

    total_reg = db.query(Registration).filter(Registration.event_id == event_id).count()
    confirmed_reg = db.query(Registration).filter(
        Registration.event_id == event_id, Registration.status == "confirmed"
    ).count()
    waitlisted_reg = db.query(Registration).filter(
        Registration.event_id == event_id, Registration.status == "waitlisted"
    ).count()
    cancelled_reg = db.query(Registration).filter(
        Registration.event_id == event_id, Registration.status == "cancelled"
    ).count()
    pending_reg = db.query(Registration).filter(
        Registration.event_id == event_id, Registration.status == "pending_confirmation"
    ).count()

    total_checkins = db.query(CheckInRecord).filter(CheckInRecord.event_id == event_id).count()

    no_show_count = confirmed_reg - total_checkins
    if no_show_count < 0:
        no_show_count = 0
    no_show_rate = (no_show_count / confirmed_reg) if confirmed_reg > 0 else 0.0

    entrance_stats = db.query(
        CheckInRecord.entrance,
        func.count(CheckInRecord.id)
    ).filter(CheckInRecord.event_id == event_id).group_by(CheckInRecord.entrance).all()
    entrance_counts = {e: c for e, c in entrance_stats}

    checkin_time_stats = db.query(
        cast(CheckInRecord.check_in_time, Date),
        func.count(CheckInRecord.id)
    ).filter(CheckInRecord.event_id == event_id).group_by(
        cast(CheckInRecord.check_in_time, Date)
    ).order_by(cast(CheckInRecord.check_in_time, Date)).all()
    checkin_timeline = [{"date": d.isoformat() if d else "", "count": c} for d, c in checkin_time_stats]

    reg_time_stats = db.query(
        cast(Registration.created_at, Date),
        func.count(Registration.id)
    ).filter(Registration.event_id == event_id).group_by(
        cast(Registration.created_at, Date)
    ).order_by(cast(Registration.created_at, Date)).all()
    registration_timeline = [{"date": d.isoformat() if d else "", "count": c} for d, c in reg_time_stats]

    status_distribution = {
        "confirmed": confirmed_reg,
        "waitlisted": waitlisted_reg,
        "cancelled": cancelled_reg,
        "pending_confirmation": pending_reg
    }

    checkin_rate = (total_checkins / confirmed_reg) if confirmed_reg > 0 else 0.0
    registration_rate = (confirmed_reg / event.max_capacity) if event.max_capacity > 0 else 0.0

    hourly_distribution = _compute_hourly_distribution(event_id, db)
    form_field_stats = _compute_form_field_stats(event_id, db)

    return {
        "event_id": event.id,
        "event_title": event.title,
        "event_location": event.location,
        "event_start_time": event.start_time,
        "event_end_time": event.end_time,
        "max_capacity": event.max_capacity,
        "total_registrations": total_reg,
        "confirmed_registrations": confirmed_reg,
        "waitlisted_registrations": waitlisted_reg,
        "cancelled_registrations": cancelled_reg,
        "pending_confirmation_registrations": pending_reg,
        "total_checkins": total_checkins,
        "checkin_rate": checkin_rate,
        "registration_rate": registration_rate,
        "entrance_counts": entrance_counts,
        "checkin_timeline": checkin_timeline,
        "registration_timeline": registration_timeline,
        "status_distribution": status_distribution,
        "checkin_hourly_distribution": hourly_distribution,
        "form_field_stats": form_field_stats,
        "no_show_count": no_show_count,
        "no_show_rate": no_show_rate
    }

@router.get("/event/{event_id}/statistics", response_model=EventStatistics)
def get_event_statistics(
    event_id: int,
    db: Session = Depends(get_db),
    organizer: User = Depends(get_current_organizer)
):
    _get_event_and_verify(event_id, organizer, db)
    stats = _compute_statistics(event_id, db)
    if not stats:
        raise HTTPException(status_code=404, detail="Event not found")
    return stats

@router.get("/event/{event_id}/export/registrations")
def export_registrations(
    event_id: int,
    db: Session = Depends(get_db),
    organizer: User = Depends(get_current_organizer)
):
    event = _get_event_and_verify(event_id, organizer, db)
    registrations = db.query(Registration).filter(Registration.event_id == event_id).all()
    form_fields = _collect_form_fields(registrations)
    stats = _compute_statistics(event_id, db)

    wb = Workbook()
    _build_event_info_sheet(wb, event, stats)
    _build_summary_sheet(wb, event, stats, db)
    _build_registrations_sheet(wb, event, db, form_fields)
    _build_form_fields_sheet(wb, stats)

    buffer = _export_workbook_to_bytes(wb)
    date_str = datetime.utcnow().strftime('%Y%m%d')
    filename = f"活动报名数据_{event.title}_{date_str}.xlsx"
    ascii_fallback = f"registrations_event{event.id}_{date_str}.xlsx"
    return _make_export_response(buffer, filename, ascii_fallback)

@router.get("/event/{event_id}/export/checkins")
def export_checkins(
    event_id: int,
    db: Session = Depends(get_db),
    organizer: User = Depends(get_current_organizer)
):
    event = _get_event_and_verify(event_id, organizer, db)
    registrations = db.query(Registration).filter(Registration.event_id == event_id).all()
    form_fields = _collect_form_fields(registrations)
    stats = _compute_statistics(event_id, db)

    wb = Workbook()
    _build_event_info_sheet(wb, event, stats)
    _build_summary_sheet(wb, event, stats, db)
    _build_checkin_sheet(wb, event, db)
    _build_hourly_distribution_sheet(wb, stats)
    _build_no_show_sheet(wb, event, db)

    buffer = _export_workbook_to_bytes(wb)
    date_str = datetime.utcnow().strftime('%Y%m%d')
    filename = f"活动签到报告_{event.title}_{date_str}.xlsx"
    ascii_fallback = f"checkins_event{event.id}_{date_str}.xlsx"
    return _make_export_response(buffer, filename, ascii_fallback)

@router.get("/event/{event_id}/export/full")
def export_full_report(
    event_id: int,
    db: Session = Depends(get_db),
    organizer: User = Depends(get_current_organizer)
):
    event = _get_event_and_verify(event_id, organizer, db)
    registrations = db.query(Registration).filter(Registration.event_id == event_id).all()
    form_fields = _collect_form_fields(registrations)
    stats = _compute_statistics(event_id, db)

    wb = Workbook()
    _build_event_info_sheet(wb, event, stats)
    _build_summary_sheet(wb, event, stats, db)
    _build_registrations_sheet(wb, event, db, form_fields)
    _build_checkin_sheet(wb, event, db)
    _build_hourly_distribution_sheet(wb, stats)
    _build_form_fields_sheet(wb, stats)
    _build_no_show_sheet(wb, event, db)

    buffer = _export_workbook_to_bytes(wb)
    date_str = datetime.utcnow().strftime('%Y%m%d')
    filename = f"活动完整报告_{event.title}_{date_str}.xlsx"
    ascii_fallback = f"full_report_event{event.id}_{date_str}.xlsx"
    return _make_export_response(buffer, filename, ascii_fallback)
